import os



import logging
from exception import MyException
from PIL import Image
from fpdf import FPDF
from docx import Document





# ------------------ Data Conversion ------------------


async def image_to_pdf(imge_path:str,output_pdf_path:str):
    try:
        image=Image.open(imge_path)

        img=image.convert("RGB")

        img.save(output_pdf_path,"PDF")

    except Exception as e:
        raise MyException(e)
    

async def text_to_pdf(file_path:str,output_file_path:str):
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                # Replace common non-latin1 characters that fpdf (standard) can't handle
                line = line.replace('\u2018', "'").replace('\u2019', "'").replace('\u201c', '"').replace('\u201d', '"')
                pdf.cell(200, 10, txt=line.strip(), ln=1)
                
        pdf.output(output_file_path)

    except Exception as e:
        raise MyException(e)



async def txt_to_pdf(txt_path, output_pdf_path):
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        with open(txt_path, "r", encoding="utf-8") as f:
            for line in f:
                pdf.cell(200, 10, txt=line.strip(), ln=1)
                
        pdf.output(output_pdf_path)

    except Exception as e:
        raise MyException(e)    



async def docs_to_pdf(docs_path: str, output_pdf_path: str):
    try:
        doc = Document(docs_path)

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        for para in doc.paragraphs:
            line = para.text.strip()
            line = (
                line
                .replace('\u2018', "'")
                .replace('\u2019', "'")
                .replace('\u201c', '"')
                .replace('\u201d', '"')
                .replace('\u2013', '-')
                .replace('\u2014', '--')
            )
            if line:
                pdf.multi_cell(0, 8, txt=line)
            else:
                pdf.ln(4)

        pdf.output(output_pdf_path)

    except Exception as e:
        raise MyException(e)
